from pathlib import Path
from io import BytesIO
import asyncio

import httpx

from docx import Document
from fastapi.testclient import TestClient

from backend import main
from backend.tools import filesystem


def _client(monkeypatch, tmp_path):
    filesystem.set_workspace_path(str(tmp_path / "workspace"))

    async def model_status():
        return {"state": "unavailable", "name": main.DEFAULT_MODEL}

    monkeypatch.setattr(main, "_model_status", model_status)
    return TestClient(main.app)


def _text_pdf(text):
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(content)).encode()
        + b" >>\nstream\n"
        + content
        + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010} 00000 n \n".encode())
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(pdf)


def test_status_reports_model_and_features(monkeypatch, tmp_path):
    async def model_status():
        return {"state": "missing", "name": "test-model"}

    client = _client(monkeypatch, tmp_path)
    monkeypatch.setattr(main, "_model_status", model_status)
    response = client.get("/status")
    assert response.status_code == 200
    assert response.json() == {
        "backend": "ready",
        "model": {"state": "missing", "name": "test-model"},
        "features": {
            "web_research": main.WEB_ENABLED,
            "memory": main.MEMORY_ENABLED,
            "files": main.FILESYSTEM_ENABLED,
        },
    }


def test_model_status_distinguishes_unavailable_and_missing(monkeypatch):
    class FakeResponse:
        def __init__(self, payload=None, error=False):
            self.payload = payload
            self.error = error

        def raise_for_status(self):
            if self.error:
                raise httpx.ConnectError("offline")

        def json(self):
            return self.payload

    class FakeClient:
        response = FakeResponse(error=True)

        def __init__(self, **_kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return False

        async def get(self, _url):
            return self.response

    monkeypatch.setattr(main.httpx, "AsyncClient", FakeClient)
    assert asyncio.run(main._model_status())["state"] == "unavailable"
    FakeClient.response = FakeResponse({"models": [{"name": "another-model"}]})
    assert asyncio.run(main._model_status())["state"] == "missing"
    FakeClient.response = FakeResponse({"models": [{"name": ["malformed"]}]})
    assert asyncio.run(main._model_status())["state"] == "missing"
    FakeClient.response = FakeResponse({"models": "malformed"})
    assert asyncio.run(main._model_status())["state"] == "unavailable"


def test_chat_passes_preferences_attachment_and_messages(monkeypatch, tmp_path):
    client = _client(monkeypatch, tmp_path)
    workspace = Path(filesystem._resolve_path("notes.docx"))
    document = Document()
    document.add_paragraph("Keep the launch checklist short.")
    document.save(workspace)
    captured = {}

    async def fake_agent(**kwargs):
        captured.update(kwargs)
        return {"role": "assistant", "content": "Done."}

    monkeypatch.setattr(main, "run_agent_loop", fake_agent)
    response = client.post(
        "/chat",
        json={
            "messages": [{"role": "user", "content": "Summarize this"}],
            "context": {"research": True, "memory": True},
            "attachments": ["notes.docx"],
        },
    )
    assert response.status_code == 200
    assert response.json()["message"]["content"] == "Done."
    assert captured["messages"] == [{"role": "user", "content": "Summarize this"}]
    assert captured["research"] is True
    assert captured["recall"] is True
    assert "launch checklist" in captured["attachment_context"]


def test_chat_rejects_invalid_history_and_step_limits(monkeypatch, tmp_path):
    client = _client(monkeypatch, tmp_path)
    assert (
        client.post(
            "/chat", json={"messages": [{"role": "system", "content": "bad"}]}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/chat",
            json={
                "messages": [{"role": "user", "content": "hello"}],
                "max_steps": None,
            },
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/chat",
            json={"messages": [{"role": "user", "content": "hello"}], "max_steps": 11},
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/chat", json={"messages": [{"role": "user", "content": " "}]}
        ).status_code
        == 400
    )


def test_chat_returns_actionable_upstream_error(monkeypatch, tmp_path):
    client = _client(monkeypatch, tmp_path)

    async def fail_agent(**_kwargs):
        raise main.AgentServiceError("Start Ollama and retry.")

    monkeypatch.setattr(main, "run_agent_loop", fail_agent)
    response = client.post(
        "/chat", json={"messages": [{"role": "user", "content": "Hello"}]}
    )
    assert response.status_code == 503
    assert response.json()["detail"] == "Start Ollama and retry."


def test_chat_rejects_traversal_and_multiple_attachments(monkeypatch, tmp_path):
    client = _client(monkeypatch, tmp_path)
    base = {"messages": [{"role": "user", "content": "Read this"}]}
    assert (
        client.post(
            "/chat", json={**base, "attachments": ["../secret.txt"]}
        ).status_code
        == 400
    )
    assert (
        client.post(
            "/chat", json={**base, "attachments": ["a.txt", "b.txt"]}
        ).status_code
        == 422
    )


def test_upload_supports_docx_rejects_unsupported_and_size(monkeypatch, tmp_path):
    client = _client(monkeypatch, tmp_path)
    document = Document()
    document.add_paragraph("Readable document")
    buffer = BytesIO()
    document.save(buffer)
    response = client.post("/files", files={"file": ("brief.docx", buffer.getvalue())})
    assert response.status_code == 200
    assert "Readable document" in filesystem.read_file("brief.docx")

    unsupported = client.post("/files", files={"file": ("scan.png", b"image")})
    assert unsupported.status_code == 415

    monkeypatch.setattr(main, "MAX_FILE_BYTES", 5)
    oversized = client.post("/files", files={"file": ("large.txt", b"123456")})
    assert oversized.status_code == 413
    assert not Path(filesystem._resolve_path("large.txt")).exists()


def test_upload_and_extract_text_pdf_and_reject_scanned_pdf(monkeypatch, tmp_path):
    client = _client(monkeypatch, tmp_path)
    response = client.post(
        "/files", files={"file": ("brief.pdf", _text_pdf("PDF content"))}
    )
    assert response.status_code == 200
    assert "PDF content" in filesystem.read_file("brief.pdf")

    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    scanned = BytesIO()
    writer.write(scanned)
    response = client.post("/files", files={"file": ("scan.pdf", scanned.getvalue())})
    assert response.status_code == 400
    assert "Scanned PDFs" in response.json()["detail"]


def test_upload_rejects_corrupt_document(monkeypatch, tmp_path):
    client = _client(monkeypatch, tmp_path)
    response = client.post("/files", files={"file": ("bad.docx", b"not a docx")})
    assert response.status_code == 400
    assert not Path(filesystem._resolve_path("bad.docx")).exists()


def test_text_extraction_is_bounded(monkeypatch, tmp_path):
    client = _client(monkeypatch, tmp_path)
    content = b"a" * (filesystem.MAX_EXTRACTED_CHARS + 5)
    uploaded = client.post("/files", files={"file": ("long.txt", content)})
    assert uploaded.status_code == 200
    extracted = filesystem.read_file("long.txt")
    assert "Document excerpt truncated" in extracted
    assert len(extracted) < filesystem.MAX_EXTRACTED_CHARS + 100


def test_desktop_token_is_required(monkeypatch, tmp_path):
    monkeypatch.setenv("ATLAS_API_TOKEN", "secret")
    client = _client(monkeypatch, tmp_path)
    assert client.get("/status").status_code == 401
    assert client.get("/health").status_code == 401
    assert client.get("/status", headers={"X-Atlas-Token": "secret"}).status_code == 200
