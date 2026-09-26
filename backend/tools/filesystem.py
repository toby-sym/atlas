"""Safe workspace file access and local document text extraction."""

import os
import shutil
from functools import lru_cache
from pathlib import Path
from uuid import UUID

import numpy as np
import pypdfium2 as pdfium
from PIL import Image
from rapidocr import RapidOCR

from backend.agent import register_tool
from backend.projects import GENERAL_PROJECT_ID, project_exists

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 30_000
TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".csv",
    ".tsv",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".html",
    ".htm",
    ".css",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".py",
    ".rs",
    ".go",
    ".java",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".sh",
    ".ps1",
    ".toml",
    ".ini",
    ".log",
    ".sql",
    ".env",
    ".diff",
    ".patch",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | IMAGE_SUFFIXES | {".pdf", ".docx"}
MAX_OCR_PAGES = 10
SAFE_ROOT = os.path.abspath(os.getenv("ATLAS_WORKSPACE", "./workspace"))


def set_workspace_path(path: str) -> None:
    # This is initialized once by the API from loaded desktop configuration.
    # pylint: disable=global-statement
    global SAFE_ROOT
    SAFE_ROOT = os.path.abspath(path)


def _resolve_path(relative_path: str) -> str:
    os.makedirs(SAFE_ROOT, exist_ok=True)
    root = os.path.realpath(SAFE_ROOT)
    target = os.path.realpath(os.path.join(root, relative_path))
    if os.path.commonpath([root, target]) != root:
        raise ValueError("Access outside safe workspace directory is prohibited.")
    return target


def _project_root(project_id: str) -> str:
    if project_id == GENERAL_PROJECT_ID:
        return _resolve_path(".")
    try:
        normalized_id = str(UUID(project_id))
    except ValueError as exc:
        raise ValueError("Invalid project ID.") from exc
    if normalized_id != project_id or not project_exists(project_id):
        raise ValueError("Project not found.")
    projects_directory = Path(os.path.realpath(SAFE_ROOT)) / ".atlas-projects"
    project_directory = projects_directory / project_id
    if projects_directory.is_symlink() or project_directory.is_symlink():
        raise ValueError("Project folders cannot be symbolic links.")
    return _resolve_path(os.path.join(".atlas-projects", project_id))


def _resolve_project_path(relative_path: str, project_id: str) -> str:
    root = os.path.realpath(_project_root(project_id))
    os.makedirs(root, exist_ok=True)
    target = os.path.realpath(os.path.join(root, relative_path))
    if os.path.commonpath([root, target]) != root:
        raise ValueError("Access outside this project directory is prohibited.")
    return target


def move_project_files_to_general(project_id: str) -> None:
    """Preserve a project's files by moving them into the General workspace."""
    if project_id == GENERAL_PROJECT_ID:
        return
    source_root = Path(_project_root(project_id))
    if not source_root.exists():
        return
    general_root = Path(_resolve_project_path(".", GENERAL_PROJECT_ID))
    moved: list[tuple[Path, Path]] = []
    try:
        entries = sorted(source_root.rglob("*"), key=lambda path: len(path.parts))
        for source in entries:
            if source.is_symlink():
                raise ValueError("Project links must be removed before deleting it.")
            if not source.is_file():
                continue
            target = general_root / source.name
            suffix = 2
            while target.exists():
                target = general_root / f"{source.stem}-{suffix}{source.suffix}"
                suffix += 1
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(source), str(target))
            moved.append((target, source))
        for directory in sorted(
            (entry for entry in entries if entry.is_dir()),
            key=lambda path: len(path.parts),
            reverse=True,
        ):
            directory.rmdir()
        source_root.rmdir()
    except Exception:
        for target, source in reversed(moved):
            source.parent.mkdir(parents=True, exist_ok=True)
            if target.exists():
                shutil.move(str(target), str(source))
        raise


@lru_cache(maxsize=1)
def _ocr_engine() -> RapidOCR:
    return RapidOCR()


def _ocr_image(image: Image.Image) -> str:
    result = _ocr_engine()(np.asarray(image.convert("RGB")))
    return "\n".join(result.txts or ())


def _ocr_pdf_page(document: pdfium.PdfDocument, index: int) -> str:
    page = document[index]
    try:
        width, height = page.get_size()
        scale = min(2.0, 2000 / max(width, height))
        bitmap = page.render(scale=scale)
        try:
            return _ocr_image(bitmap.to_pil())
        finally:
            bitmap.close()
    finally:
        page.close()


def extract_file(path: str, project_id: str = GENERAL_PROJECT_ID) -> str:  # pylint: disable=import-outside-toplevel,too-many-branches
    """Read a supported workspace document as text, with bounded extraction."""
    resolved = _resolve_project_path(path, project_id)
    if os.path.getsize(resolved) > MAX_FILE_BYTES:
        raise ValueError("Files must be 10 MB or smaller.")
    suffix = Path(resolved).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(
            "This file type is not supported. Use text, code, PDF, DOCX, PNG, or JPEG."
        )

    if suffix in TEXT_SUFFIXES:
        content = Path(resolved).read_text(encoding="utf-8-sig")
    elif suffix in IMAGE_SUFFIXES:
        with Image.open(resolved) as image:
            content = _ocr_image(image).strip()
        if not content:
            raise ValueError("This image has no text that OCR could recognize.")
    elif suffix == ".pdf":
        from pypdf import PdfReader  # pylint: disable=import-outside-toplevel

        reader = PdfReader(resolved, strict=True)
        page_texts = [page.extract_text() or "" for page in reader.pages]
        ocr_limited = False
        if any(not text.strip() for text in page_texts):
            document = pdfium.PdfDocument(resolved)
            try:
                ocr_count = 0
                for index, text in enumerate(page_texts):
                    if not text.strip() and ocr_count < MAX_OCR_PAGES:
                        page_texts[index] = _ocr_pdf_page(document, index)
                        ocr_count += 1
                ocr_limited = (
                    any(not text.strip() for text in page_texts)
                    and ocr_count >= MAX_OCR_PAGES
                )
            finally:
                document.close()
        content = "\n\n".join(page_texts).strip()
        if not content:
            raise ValueError(
                "This PDF has no text that extraction or OCR could recognize."
            )
        if ocr_limited:
            content += "\n\n[OCR stopped after 10 image pages.]"
    else:
        from docx import Document  # pylint: disable=import-outside-toplevel

        document = Document(resolved)
        paragraphs = [
            paragraph.text for paragraph in document.paragraphs if paragraph.text
        ]
        for table in document.tables:
            paragraphs.extend(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
        content = "\n".join(paragraphs).strip()
        if not content:
            raise ValueError("This DOCX document contains no extractable text.")

    if len(content) > MAX_EXTRACTED_CHARS:
        return (
            content[:MAX_EXTRACTED_CHARS]
            + "\n\n[Document excerpt truncated at 30,000 characters.]"
        )
    return content


@register_tool(
    name="read_file",
    description="Read text, code, PDF, DOCX, or OCR image content from the local workspace.",
    parameters={
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "Relative file path inside workspace",
            },
        },
        "required": ["path"],
    },
)
def read_file(path: str, project_id: str = GENERAL_PROJECT_ID) -> str:
    try:
        return extract_file(path, project_id)
    except Exception as exc:  # pylint: disable=broad-exception-caught
        return f"Error reading file: {exc}"
